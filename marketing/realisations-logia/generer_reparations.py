"""Génère l'exemple de réalisation Logia « Gestion de réparations de téléphones » (Excel, compatible Google Sheets).

Modèle de démonstration : ni offert ni publié.
Usage : python3 generer_reparations.py  ->  Gestion-reparations-telephones-DWTS.xlsx
"""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from base_logia import *

SORTIE = Path(__file__).with_name('Gestion-reparations-telephones-DWTS.xlsx')
NB_CLIENTS, NB_REP = 150, 400
STATUTS = ['En attente', 'En cours', 'Terminé', 'Livré', 'Abandonné']
COULEURS = {'En attente': ('FEF3C7', '92400E'), 'En cours': ('DBEAFE', '1E40AF'), 'Terminé': ('D1F7E0', '166534'),
            'Livré': ('E7EAF0', '555555'), 'Abandonné': ('FEE2E2', 'B91C1C')}

wb = Workbook()
menu(wb, 'Gestion de réparations de téléphones', [
    ('Clients', 'Une ligne par client : nom et téléphone. Son numéro client (C001…) se crée seul.'),
    ('Réparations', "Chaque téléphone déposé : panne, devis, acompte, état. Le reste à payer se calcule."),
    ('Tableau de bord', "Combien de téléphones par état, chiffre d'affaires, argent restant à encaisser."),
    ('Paramètres', "Nom de l'atelier et devise."),
])
pa, (ATELIER, DEVISE) = parametres(wb, [("Nom de l'atelier", 'Mon atelier'), ('Devise', 'XAF')])

# ---------- Clients ----------
cl = wb.create_sheet('Clients', 1)
bandeau(cl, 'Clients', 4)
retour_menu(cl, 1)
entetes(cl, 4, ['N° client', 'Nom et prénom', 'Téléphone', 'Nombre de réparations'], [12, 32, 18, 22])
for i in range(NB_CLIENTS):
    r = 5 + i
    cl.cell(r, 1, f'=IF(B{r}="","","C"&TEXT(ROW()-4,"000"))')
    cl.cell(r, 4, f'=IF(B{r}="","",COUNTIF(Réparations!B5:B{4 + NB_REP},A{r}))')
for i, (n, t) in enumerate([('Patrick Moussavou', '066 12 34 56'), ('Sandrine Mba', '077 45 67 89'), ('Armel Ndong', '062 98 76 54')]):
    cl.cell(5 + i, 2, n); cl.cell(5 + i, 3, t)
corps(cl, 5, NB_CLIENTS, 4, calcule=(1, 4), centre=(1, 4))
cl.freeze_panes = 'A5'

# ---------- Réparations ----------
re = wb.create_sheet('Réparations', 2)
bandeau(re, 'Réparations', 12)
re['A1'] = f'={ATELIER}&" · Réparations"'
retour_menu(re, 1)
T = ['N° dossier', 'N° client', 'Client', 'Date de dépôt', 'Téléphone (marque et modèle)', 'Panne', 'Prix du devis',
     'Acompte versé', 'Reste à payer', 'État', 'Date de remise', 'Remarques']
entetes(re, 4, T, [11, 11, 24, 14, 26, 28, 14, 14, 14, 13, 14, 26])
exemples = [
    ('C001', date(2026, 10, 20), 'Samsung Galaxy A14', 'Écran cassé', 45000, 20000, 'En cours', None),
    ('C002', date(2026, 10, 22), 'iPhone 11', 'Batterie usée', 30000, 30000, 'Terminé', None),
    ('C003', date(2026, 10, 10), 'Tecno Spark 10', 'Ne charge plus', 12000, 5000, 'Livré', date(2026, 10, 14)),
    ('C001', date(2026, 10, 26), 'Infinix Hot 30', "Problème d'allumage", 25000, 0, 'En attente', None),
]
for i in range(NB_REP):
    r = 5 + i
    re.cell(r, 1, f'=IF(B{r}="","","R"&TEXT(ROW()-4,"0000"))')
    re.cell(r, 3, f'=IF(B{r}="","",IFERROR(INDEX(Clients!B$5:B${4 + NB_CLIENTS},MATCH(B{r},Clients!A$5:A${4 + NB_CLIENTS},0)),"Client inconnu"))')
    re.cell(r, 9, f'=IF(G{r}="","",G{r}-N(H{r}))')
    if i < len(exemples):
        c, d, tel, panne, devis, ac, st, rem = exemples[i]
        for col, v in [(2, c), (4, d), (5, tel), (6, panne), (7, devis), (8, ac), (10, st), (11, rem)]:
            re.cell(r, col, v)
    for col in (4, 11):
        re.cell(r, col).number_format = 'dd/mm/yyyy'
    for col in (7, 8, 9):
        re.cell(r, col).number_format = '# ##0'
corps(re, 5, NB_REP, 12, calcule=(1, 3, 9), centre=(1, 2, 4, 10, 11))
dv = DataValidation(type='list', formula1='"' + ','.join(STATUTS) + '"', allow_blank=True)
re.add_data_validation(dv); dv.add(f'J5:J{4 + NB_REP}')
for s, (bg, fg) in COULEURS.items():
    re.conditional_formatting.add(f'J5:J{4 + NB_REP}', FormulaRule(formula=[f'J5="{s}"'], fill=fond(bg), font=Font(bold=True, color=fg)))
re.conditional_formatting.add(f'I5:I{4 + NB_REP}', FormulaRule(formula=['AND(ISNUMBER(I5),I5>0,J5<>"Abandonné")'], font=Font(bold=True, color='B91C1C')))
re.auto_filter.ref = f'A4:L{4 + NB_REP}'
re.freeze_panes = 'D5'

# ---------- Tableau de bord ----------
tb = wb.create_sheet('Tableau de bord', 3)
bandeau(tb, 'Tableau de bord', 6)
retour_menu(tb, 1)
for c in 'ABCDEF':
    tb.column_dimensions[c].width = 22
plage = lambda col: f'Réparations!{col}5:{col}{4 + NB_REP}'
carte(tb, 4, 1, 'Dossiers au total', f'=COUNTA({plage("B")})')
carte(tb, 4, 2, 'Téléphones à rendre', f'=COUNTIF({plage("J")},"En attente")+COUNTIF({plage("J")},"En cours")+COUNTIF({plage("J")},"Terminé")')
carte(tb, 4, 3, "Chiffre d'affaires (devis hors abandonnés)", f'=SUM({plage("G")})-SUMIF({plage("J")},"Abandonné",{plage("G")})', '# ##0')
carte(tb, 4, 4, 'Acomptes encaissés', f'=SUM({plage("H")})', '# ##0')
carte(tb, 4, 5, 'Reste à encaisser', f'=SUMIF({plage("J")},"<>Abandonné",{plage("I")})', '# ##0')
for col, t in [(1, 'État'), (2, 'Dossiers')]:
    c = tb.cell(8, col, t)
    c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = fond(BLEU); c.alignment = CENTRE; c.border = CADRE
for i, s in enumerate(STATUTS):
    r = 9 + i
    tb.cell(r, 1, s)
    tb.cell(r, 2, f'=COUNTIF({plage("J")},A{r})')
corps(tb, 9, len(STATUTS), 2, calcule=(2,), centre=(2,))

wb.save(SORTIE)
print(SORTIE)
