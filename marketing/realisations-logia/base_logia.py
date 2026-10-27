"""Éléments communs aux exemples de réalisation Logia (couleurs, en-têtes, menu, tableaux)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.hyperlink import Hyperlink

NAVY, BLEU, CYAN, GRIS = '0F2A4A', '1E6FD9', '22D3EE', 'F3F5F9'
SAISIE = 'FFF7D6'
fin = Side(style='thin', color='D5DBE5')
CADRE = Border(left=fin, right=fin, top=fin, bottom=fin)
CENTRE = Alignment(horizontal='center', vertical='center')
GAUCHE = Alignment(horizontal='left', vertical='center', indent=1)
SIGNATURE = 'Réalisé par Logia (DWTS) · Un logiciel sur mesure pour votre entreprise ? WhatsApp 062 67 14 45'


def fond(c):
    return PatternFill('solid', start_color=c, end_color=c)


def bandeau(ws, titre, largeur, sous_titre=SIGNATURE):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=largeur)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=largeur)
    ws['A1'] = titre
    ws['A1'].font = Font(name='Arial', size=18, bold=True, color='FFFFFF')
    ws['A2'] = sous_titre
    ws['A2'].font = Font(name='Arial', size=10, color='CFE3FF')
    for r in (1, 2):
        for c in range(1, largeur + 1):
            ws.cell(r, c).fill = fond(NAVY)
        ws.cell(r, 1).alignment = GAUCHE
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.sheet_view.showGridLines = False


def entetes(ws, ligne, titres, largeurs):
    for col, (t, w) in enumerate(zip(titres, largeurs), start=1):
        c = ws.cell(ligne, col, t)
        c.font = Font(name='Arial', bold=True, color='FFFFFF')
        c.fill = fond(BLEU)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = CADRE
        ws.column_dimensions[L(col)].width = w
    ws.row_dimensions[ligne].height = 30


def corps(ws, ligne0, nb, nb_col, calcule=(), centre=()):
    """Bordures, police et lignes alternées ; `calcule` = colonnes à formule (fond bleu clair, pas de saisie)."""
    for i in range(nb):
        for col in range(1, nb_col + 1):
            c = ws.cell(ligne0 + i, col)
            c.border = CADRE
            c.font = Font(name='Arial', bold=col in calcule)
            c.alignment = CENTRE if col in centre else GAUCHE
            if col in calcule:
                c.fill = fond('EEF4FF')
            elif i % 2:
                c.fill = fond(GRIS)
        ws.row_dimensions[ligne0 + i].height = 20


def lien(cell, feuille, texte=None):
    if texte:
        cell.value = texte
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{feuille}'!A1")


def menu(wb, titre, boutons):
    """Feuille « Menu principal » : un gros bouton-lien par feuille (aucune macro, marche aussi sur Google Sheets)."""
    ws = wb.active
    ws.title = 'Menu principal'
    bandeau(ws, titre, 5)
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 62
    r = 4
    for feuille, aide in boutons:
        b = ws.cell(r, 2)
        lien(b, feuille, f'  ▶  {feuille}')
        b.font = Font(name='Arial', size=13, bold=True, color='FFFFFF', underline='none')
        b.fill = fond(BLEU)
        b.alignment = GAUCHE
        ws.cell(r, 3, aide).font = Font(name='Arial', color='555555')
        ws.cell(r, 3).alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 34
        r += 2
    ws.cell(r, 2, 'Cases jaunes = à remplir · cases bleu clair = calcul automatique, ne pas écrire dedans.').font = Font(name='Arial', italic=True, color='555555')
    return ws


def retour_menu(ws, col):
    lien(ws.cell(3, col), 'Menu principal', '◀ Menu principal')
    ws.cell(3, col).font = Font(name='Arial', bold=True, color=BLEU)
    ws.cell(3, col).hyperlink.location = "'Menu principal'!A1"


def parametres(wb, lignes):
    """Feuille « Paramètres » : liste de (libellé, valeur). Retourne les références absolues dans l'ordre."""
    pa = wb.create_sheet('Paramètres')
    bandeau(pa, 'Paramètres', 4)
    pa.column_dimensions['A'].width = 30
    pa.column_dimensions['B'].width = 34
    refs = []
    for i, (lib, val) in enumerate(lignes, start=4):
        pa.cell(i, 1, lib).font = Font(name='Arial', bold=True, color=NAVY)
        v = pa.cell(i, 2, val)
        v.font = Font(name='Arial', size=12, bold=True)
        v.fill = fond(SAISIE)
        v.border = CADRE
        refs.append(f"'Paramètres'!$B${i}")
    return pa, refs


def carte(ws, ligne, col, titre, formule, fmt='0'):
    """Petite carte de chiffre clé (titre au-dessus, valeur dessous)."""
    t = ws.cell(ligne, col, titre)
    t.font = Font(name='Arial', size=9, bold=True, color='FFFFFF')
    t.fill = fond(NAVY)
    t.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    v = ws.cell(ligne + 1, col, formule)
    v.font = Font(name='Arial', size=16, bold=True, color=NAVY)
    v.fill = fond('EEF4FF')
    v.alignment = CENTRE
    v.number_format = fmt
    for c in (t, v):
        c.border = CADRE
    ws.row_dimensions[ligne].height = 30
    ws.row_dimensions[ligne + 1].height = 36
