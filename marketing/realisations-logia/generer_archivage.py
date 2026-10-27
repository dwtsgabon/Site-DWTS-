"""Génère l'exemple de réalisation Logia « Archivage électronique des documents » (Excel, compatible Google Sheets).

Modèle de démonstration : ni offert ni publié.
Usage : python3 generer_archivage.py  ->  Archivage-des-documents-DWTS.xlsx
"""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from base_logia import *

SORTIE = Path(__file__).with_name('Archivage-des-documents-DWTS.xlsx')
NB = 300
CATEGORIES = ['Facture', 'Contrat', 'Courrier', 'Rapport', 'Administratif', 'Ressources humaines', 'Autre']
STATUTS = ['Actif', 'À classer', 'Archivé']

wb = Workbook()
menu(wb, 'Archivage électronique des documents', [
    ('Documents', "Enregistrez chaque document : titre, catégorie, lien vers le fichier. Filtrez et cherchez dans les colonnes."),
    ('Tableau de bord', 'Nombre de documents par catégorie et par statut, sans rien calculer.'),
    ('Paramètres', "Nom de l'entreprise et listes des catégories."),
])

pa, (ENT,) = parametres(wb, [("Nom de l'entreprise", 'Mon entreprise')])
pa.cell(6, 1, 'Catégories').font = Font(name='Arial', bold=True, color=NAVY)
for i, cat in enumerate(CATEGORIES):
    c = pa.cell(7 + i, 1, cat)
    c.fill = fond(SAISIE)
    c.border = CADRE
    c.font = Font(name='Arial')
CAT_PLAGE = f"'Paramètres'!$A$7:$A${6 + len(CATEGORIES)}"

# ---------- Documents ----------
do = wb.create_sheet('Documents', 1)
bandeau(do, 'Documents', 8)
do['A1'] = f'={ENT}&" · Archive des documents"'
retour_menu(do, 1)
TITRES = ['N°', 'Date du document', 'Titre', 'Catégorie', 'Émetteur / Destinataire', 'Emplacement physique', 'Lien du fichier', 'Statut']
entetes(do, 4, TITRES, [9, 15, 38, 20, 26, 22, 34, 12])
exemples = [
    (date(2026, 1, 12), 'Facture fournisseur n°2026-014', 'Facture', 'Sogapro', 'Classeur A · pochette 3', 'Actif'),
    (date(2026, 2, 3), 'Contrat de bail du magasin', 'Contrat', 'M. Ondo (bailleur)', 'Armoire 1', 'Archivé'),
    (date(2026, 3, 18), 'Courrier de la mairie', 'Courrier', 'Mairie de Port-Gentil', 'Classeur B', 'À classer'),
    (date(2026, 4, 2), 'Rapport de ventes du 1er trimestre', 'Rapport', 'Direction', 'Numérique seulement', 'Actif'),
]
for i in range(NB):
    r = 5 + i
    do.cell(r, 1, f'=IF(C{r}="","","D"&TEXT(ROW()-4,"0000"))')
    if i < len(exemples):
        d, t, cat, em, ep, st = exemples[i]
        do.cell(r, 2, d); do.cell(r, 3, t); do.cell(r, 4, cat); do.cell(r, 5, em); do.cell(r, 6, ep); do.cell(r, 8, st)
    do.cell(r, 2).number_format = 'dd/mm/yyyy'
corps(do, 5, NB, 8, calcule=(1,), centre=(1, 2, 8))
do.cell(5, 7, 'https://exemple.com/facture-2026-014.pdf')
do.cell(5, 7).hyperlink = 'https://exemple.com/facture-2026-014.pdf'
do.cell(5, 7).font = Font(name='Arial', color=BLEU, underline='single')
dv = DataValidation(type='list', formula1=f'={CAT_PLAGE}', allow_blank=True)
dv2 = DataValidation(type='list', formula1='"' + ','.join(STATUTS) + '"', allow_blank=True)
do.add_data_validation(dv); do.add_data_validation(dv2)
dv.add(f'D5:D{4 + NB}'); dv2.add(f'H5:H{4 + NB}')
do.conditional_formatting.add(f'H5:H{4 + NB}', FormulaRule(formula=['H5="À classer"'], fill=fond('FEF3C7'), font=Font(bold=True, color='92400E')))
do.conditional_formatting.add(f'H5:H{4 + NB}', FormulaRule(formula=['H5="Archivé"'], fill=fond('E7EAF0'), font=Font(color='555555')))
do.auto_filter.ref = f'A4:H{4 + NB}'
do.freeze_panes = 'A5'

# ---------- Tableau de bord ----------
tb = wb.create_sheet('Tableau de bord', 2)
bandeau(tb, 'Tableau de bord', 6)
retour_menu(tb, 1)
for c, w in zip('ABCDEF', [26, 14, 4, 20, 14, 4]):
    tb.column_dimensions[c].width = w
carte(tb, 4, 1, 'Documents au total', f'=COUNTA(Documents!C5:C{4 + NB})')
carte(tb, 4, 2, 'À classer', f'=COUNTIF(Documents!H5:H{4 + NB},"À classer")')
tb.cell(8, 1, 'Catégorie').font = Font(name='Arial', bold=True, color='FFFFFF')
tb.cell(8, 2, 'Documents').font = Font(name='Arial', bold=True, color='FFFFFF')
for col in (1, 2):
    tb.cell(8, col).fill = fond(BLEU); tb.cell(8, col).alignment = CENTRE; tb.cell(8, col).border = CADRE
for i in range(len(CATEGORIES)):
    r = 9 + i
    tb.cell(r, 1, f"='Paramètres'!A{7 + i}")
    tb.cell(r, 2, f'=IF(A{r}="","",COUNTIF(Documents!D5:D{4 + NB},A{r}))')
corps(tb, 9, len(CATEGORIES), 2, centre=(2,))
tb.cell(8, 4, 'Statut').font = Font(name='Arial', bold=True, color='FFFFFF')
tb.cell(8, 5, 'Documents').font = Font(name='Arial', bold=True, color='FFFFFF')
for col in (4, 5):
    tb.cell(8, col).fill = fond(BLEU); tb.cell(8, col).alignment = CENTRE; tb.cell(8, col).border = CADRE
for i, s in enumerate(STATUTS):
    r = 9 + i
    tb.cell(r, 4, s)
    tb.cell(r, 5, f'=COUNTIF(Documents!H5:H{4 + NB},D{r})')
    for col in (4, 5):
        tb.cell(r, col).border = CADRE
        tb.cell(r, col).font = Font(name='Arial', bold=col == 5)
        tb.cell(r, col).alignment = CENTRE if col == 5 else GAUCHE

wb.save(SORTIE)
print(SORTIE)
