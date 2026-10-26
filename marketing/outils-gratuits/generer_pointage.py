"""Génère l'outil gratuit « Pointage du personnel » (Excel, compatible Google Sheets).

Usage : python3 generer_pointage.py  ->  Pointage-du-personnel-DWTS.xlsx
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

SORTIE = Path(__file__).with_name('Pointage-du-personnel-DWTS.xlsx')
NB_EMPLOYES = 30
MOIS = ['Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet',
        'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
CODES = [('P', 'Présent', 'D1F7E0', '166534'), ('R', 'En retard', 'FEF3C7', '92400E'),
         ('A', 'Absent', 'FEE2E2', 'B91C1C'), ('C', 'Congé', 'DBEAFE', '1E40AF'),
         ('M', 'Malade', 'EDE9FE', '6D28D9')]

NAVY, BLEU, CYAN, GRIS = '0F2A4A', '1E6FD9', '22D3EE', 'F3F5F9'
fin = Side(style='thin', color='D5DBE5')
CADRE = Border(left=fin, right=fin, top=fin, bottom=fin)
CENTRE = Alignment(horizontal='center', vertical='center')
GAUCHE = Alignment(horizontal='left', vertical='center', indent=1)


def fond(c):
    return PatternFill('solid', start_color=c, end_color=c)


def bandeau(ws, titre, sous_titre, largeur):
    """En-tête de marque commun à toutes les feuilles."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=largeur)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=largeur)
    ws['A1'] = titre
    ws['A1'].font = Font(name='Arial', size=18, bold=True, color='FFFFFF')
    ws['A2'] = sous_titre
    ws['A2'].font = Font(name='Arial', size=10, color='CFE3FF')
    for r in (1, 2):
        for c in range(1, largeur + 1):
            ws.cell(r, c).fill = fond(NAVY)
        ws.cell(r, 1).alignment = GAUCHE
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20


SIGNATURE = 'Outil offert par Logia (DWTS) · Un logiciel sur mesure pour votre entreprise ? WhatsApp 062 67 14 45'

wb = Workbook()

# ---------- Mode d'emploi ----------
ws = wb.active
ws.title = "Mode d'emploi"
bandeau(ws, 'Pointage du personnel', SIGNATURE, 6)
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 70
etapes = [
    ('1. Paramètres', "Écrivez le nom de votre entreprise et l'année."),
    ('2. Employés', 'Listez vos employés (jusqu\'à 30) et, si vous voulez, leur salaire par jour.'),
    ('3. Pointage', 'Chaque jour, ouvrez le mois et choisissez un code dans la case : P, R, A, C ou M.'),
    ('4. Totaux', 'Les présences, absences, taux de présence et salaire dû se calculent seuls.'),
    ('5. Bilan annuel', "Le total de l'année par employé, sans rien faire."),
]
r = 4
for b, c in etapes:
    ws.cell(r, 2, b).font = Font(name='Arial', bold=True, color=BLEU)
    ws.cell(r, 3, c).font = Font(name='Arial')
    r += 1
r += 1
ws.cell(r, 2, 'Codes').font = Font(name='Arial', bold=True, color=NAVY)
r += 1
for code, nom, bg, fg in CODES:
    x = ws.cell(r, 2, code)
    x.font = Font(name='Arial', bold=True, color=fg)
    x.fill = fond(bg)
    x.alignment = CENTRE
    ws.cell(r, 3, nom).font = Font(name='Arial')
    r += 1
r += 1
ws.cell(r, 2, 'Salaire dû = (jours présents + en retard + congés) × salaire par jour. Les absences et jours de maladie ne sont pas payés : modifiez la formule si votre règle est différente.').font = Font(name='Arial', italic=True, color='555555')
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[r].height = 32

# ---------- Paramètres ----------
pa = wb.create_sheet('Paramètres')
bandeau(pa, 'Paramètres', SIGNATURE, 4)
pa.column_dimensions['A'].width = 26
pa.column_dimensions['B'].width = 34
for i, (lib, val) in enumerate([("Nom de l'entreprise", 'Mon entreprise'), ('Année', 2026), ('Devise', 'XAF')], start=4):
    pa.cell(i, 1, lib).font = Font(name='Arial', bold=True, color=NAVY)
    v = pa.cell(i, 2, val)
    v.font = Font(name='Arial', size=12, bold=True)
    v.fill = fond('FFF7D6')
    v.border = CADRE
ENTREPRISE, ANNEE, DEVISE = "'Paramètres'!$B$4", "'Paramètres'!$B$5", "'Paramètres'!$B$6"

# ---------- Employés ----------
em = wb.create_sheet('Employés')
bandeau(em, 'Employés', SIGNATURE, 4)
for col, (t, w) in enumerate([('ID', 8), ('Nom et prénom', 30), ('Poste', 22), ('Salaire par jour', 18)], start=1):
    c = em.cell(4, col, t)
    c.font = Font(name='Arial', bold=True, color='FFFFFF')
    c.fill = fond(BLEU)
    c.alignment = CENTRE
    em.column_dimensions[L(col)].width = w
exemples = [('Jean Paul', 'Vendeur', 5000), ('Marie Nzé', 'Caissière', 5000)]
for i in range(NB_EMPLOYES):
    r = 5 + i
    em.cell(r, 1, f'E{i + 1:03d}').alignment = CENTRE
    if i < len(exemples):
        em.cell(r, 2, exemples[i][0]); em.cell(r, 3, exemples[i][1]); em.cell(r, 4, exemples[i][2])
    em.cell(r, 4).number_format = '# ##0'
    for col in range(1, 5):
        em.cell(r, col).border = CADRE
        em.cell(r, col).font = Font(name='Arial')
        if i % 2:
            em.cell(r, col).fill = fond(GRIS)
em.freeze_panes = 'A5'

# ---------- Feuilles mensuelles ----------
J0 = 4                      # colonne du 1er jour (D)
JF = J0 + 30                # colonne du 31
TOT = ['P', 'R', 'A', 'C', 'M']
T0 = JF + 1                 # 1re colonne de totaux
C_TAUX, C_SAL = T0 + len(TOT), T0 + len(TOT) + 1
DERN = C_SAL
LIG_DATE, LIG_JOUR, LIG0 = 5, 6, 7
LIGF = LIG0 + NB_EMPLOYES - 1

validation_codes = ','.join(c for c, *_ in CODES)
for m, nom_mois in enumerate(MOIS, start=1):
    ws = wb.create_sheet(nom_mois)
    bandeau(ws, f'Pointage · {nom_mois}', SIGNATURE, DERN)
    ws['A1'] = f'={ENTREPRISE}&" · Pointage de {nom_mois} "&{ANNEE}'
    ws.cell(3, 1, 'Codes :').font = Font(name='Arial', bold=True, color=NAVY)
    # légende compacte sur la ligne 3
    for k, (code, nom, bg, fg) in enumerate(CODES):
        c = ws.cell(3, J0 + k * 5, f'{code} = {nom}')
        c.font = Font(name='Arial', bold=True, color=fg, size=9)
        c.fill = fond(bg)
        ws.merge_cells(start_row=3, start_column=J0 + k * 5, end_row=3, end_column=J0 + k * 5 + 3)
        c.alignment = CENTRE

    for col, t in [(1, 'ID'), (2, 'Nom'), (3, 'Poste')]:
        ws.merge_cells(start_row=LIG_DATE, start_column=col, end_row=LIG_JOUR, end_column=col)
        c = ws.cell(LIG_DATE, col, t)
    for d in range(1, 32):
        col = J0 + d - 1
        date = f'DATE({ANNEE},{m},{d})'
        ws.cell(LIG_DATE, col, f'=IF(MONTH({date})={m},{d},"")')
        ws.cell(LIG_JOUR, col, f'=IF({L(col)}{LIG_DATE}="","",CHOOSE(WEEKDAY({date},2),"Lun","Mar","Mer","Jeu","Ven","Sam","Dim"))')
        ws.column_dimensions[L(col)].width = 4.6
    titres_tot = ['Présent', 'Retard', 'Absent', 'Congé', 'Malade', 'Taux de présence', 'Salaire dû']
    for k, t in enumerate(titres_tot):
        col = T0 + k
        ws.merge_cells(start_row=LIG_DATE, start_column=col, end_row=LIG_JOUR, end_column=col)
        ws.cell(LIG_DATE, col, t)
        ws.column_dimensions[L(col)].width = 9 if k < 5 else 13
    for r in (LIG_DATE, LIG_JOUR):
        for col in range(1, DERN + 1):
            c = ws.cell(r, col)
            c.font = Font(name='Arial', bold=True, color='FFFFFF', size=9)
            c.fill = fond(BLEU if col < T0 else NAVY)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = CADRE
    ws.row_dimensions[LIG_DATE].height = 22
    ws.row_dimensions[LIG_JOUR].height = 22
    ws.column_dimensions['A'].width = 7
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 14

    for i in range(NB_EMPLOYES):
        r = LIG0 + i
        e = 5 + i
        ws.cell(r, 1, f'=IF(\'Employés\'!B{e}="","",\'Employés\'!A{e})')
        ws.cell(r, 2, f'=IF(\'Employés\'!B{e}="","",\'Employés\'!B{e})')
        ws.cell(r, 3, f'=IF(\'Employés\'!C{e}="","",\'Employés\'!C{e})')
        plage = f'{L(J0)}{r}:{L(JF)}{r}'
        for k, code in enumerate(TOT):
            ws.cell(r, T0 + k, f'=IF($B{r}="","",COUNTIF({plage},"{code}"))')
        p, rr, a, cg, ml = (f'{L(T0 + k)}{r}' for k in range(5))
        ws.cell(r, C_TAUX, f'=IF($B{r}="","",IF({p}+{rr}+{a}+{ml}=0,"",({p}+{rr})/({p}+{rr}+{a}+{ml})))')
        ws.cell(r, C_TAUX).number_format = '0%'
        ws.cell(r, C_SAL, f'=IF($B{r}="","",IF(\'Employés\'!D{e}="","",({p}+{rr}+{cg})*\'Employés\'!D{e}))')
        ws.cell(r, C_SAL).number_format = '# ##0'
        for col in range(1, DERN + 1):
            c = ws.cell(r, col)
            c.border = CADRE
            c.font = Font(name='Arial', size=10, bold=col >= T0)
            c.alignment = GAUCHE if col in (2, 3) else CENTRE
            if col >= T0:
                c.fill = fond('EEF4FF')
        ws.row_dimensions[r].height = 20

    # ligne de totaux
    rt = LIGF + 1
    ws.cell(rt, 2, 'TOTAL')
    for k in range(5):
        col = T0 + k
        ws.cell(rt, col, f'=SUM({L(col)}{LIG0}:{L(col)}{LIGF})')
    ws.cell(rt, C_SAL, f'=SUM({L(C_SAL)}{LIG0}:{L(C_SAL)}{LIGF})')
    ws.cell(rt, C_SAL).number_format = '# ##0'
    for col in range(1, DERN + 1):
        c = ws.cell(rt, col)
        c.fill = fond(NAVY)
        c.font = Font(name='Arial', bold=True, color='FFFFFF')
        c.alignment = CENTRE
        c.border = CADRE

    # listes déroulantes et couleurs
    dv = DataValidation(type='list', formula1=f'"{validation_codes}"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Code inconnu',
                        error='Utilisez P, R, A, C ou M.')
    ws.add_data_validation(dv)
    zone = f'{L(J0)}{LIG0}:{L(JF)}{LIGF}'
    dv.add(zone)
    for code, nom, bg, fg in CODES:
        ws.conditional_formatting.add(zone, FormulaRule(
            formula=[f'UPPER({L(J0)}{LIG0})="{code}"'], fill=fond(bg), font=Font(bold=True, color=fg)))
    # week-ends et jours hors du mois grisés
    ws.conditional_formatting.add(zone, FormulaRule(
        formula=[f'OR({L(J0)}${LIG_JOUR}="Sam",{L(J0)}${LIG_JOUR}="Dim",{L(J0)}${LIG_DATE}="")'], fill=fond('E7EAF0')))
    ws.freeze_panes = ws.cell(LIG0, J0)
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------- Bilan annuel ----------
bi = wb.create_sheet('Bilan annuel')
bandeau(bi, 'Bilan annuel', SIGNATURE, 9)
bi['A1'] = f'={ENTREPRISE}&" · Bilan annuel "&{ANNEE}'
entetes = ['ID', 'Nom', 'Présent', 'Retard', 'Absent', 'Congé', 'Malade', 'Taux de présence', 'Salaire dû']
for col, t in enumerate(entetes, start=1):
    c = bi.cell(4, col, t)
    c.font = Font(name='Arial', bold=True, color='FFFFFF')
    c.fill = fond(BLEU)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    bi.column_dimensions[L(col)].width = [8, 26, 10, 10, 10, 10, 10, 14, 16][col - 1]
bi.row_dimensions[4].height = 30
for i in range(NB_EMPLOYES):
    r = 5 + i
    lm = LIG0 + i
    bi.cell(r, 1, f'=IF(\'Employés\'!B{r}="","",\'Employés\'!A{r})')
    bi.cell(r, 2, f'=IF(\'Employés\'!B{r}="","",\'Employés\'!B{r})')
    for k in range(5):
        col_m = L(T0 + k)
        somme = '+'.join(f"N('{mo}'!{col_m}{lm})" for mo in MOIS)
        bi.cell(r, 3 + k, f'=IF($B{r}="","",{somme})')
    bi.cell(r, 8, f'=IF($B{r}="","",IF(C{r}+D{r}+E{r}+G{r}=0,"",(C{r}+D{r})/(C{r}+D{r}+E{r}+G{r})))')
    bi.cell(r, 8).number_format = '0%'
    sal = '+'.join(f"N('{mo}'!{L(C_SAL)}{lm})" for mo in MOIS)
    bi.cell(r, 9, f'=IF($B{r}="","",{sal})')
    bi.cell(r, 9).number_format = '# ##0'
    for col in range(1, 10):
        c = bi.cell(r, col)
        c.border = CADRE
        c.font = Font(name='Arial', bold=col >= 3)
        c.alignment = GAUCHE if col == 2 else CENTRE
        if i % 2:
            c.fill = fond(GRIS)
bi.freeze_panes = 'A5'

wb.move_sheet('Bilan annuel', offset=-(len(MOIS)))
wb.save(SORTIE)
print(SORTIE)
